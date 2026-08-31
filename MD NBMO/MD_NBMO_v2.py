#------------------------------------------------ Begin_Librairie ----------------------------------------
import datetime
import os
import re
import struct
import warnings

import pandas as pd
import requests
from bs4 import BeautifulSoup
from docx import Document
from urllib.parse import urljoin, unquote

warnings.filterwarnings('ignore')
requests.packages.urllib3.disable_warnings()

#------------------------------------------------ Begin_fileName ----------------------------------------
regulatorName = 'MD NBMO'
print('Running {} Web Scraping Tool v.2'.format(regulatorName))

now = datetime.datetime.now()
processdate = now.strftime('%Y-%m-%d')
filename = '{} SQL Ready {}.xlsx'.format(regulatorName, str(now).replace(':', '.')[:-7])

try:
    scriptfolder = os.path.dirname(os.path.abspath(__file__))   # production environment (.py)
except NameError:
    scriptfolder = os.getcwd()                                  # notebook environment

os.chdir(scriptfolder)
tempfolder = os.path.join(scriptfolder, 'tempfolder')
if os.path.exists(tempfolder):
    for rem in os.listdir(tempfolder):
        os.remove(os.path.join(tempfolder, rem))
else:
    os.mkdir(tempfolder)

#------------------------------------------------ Begin_Variable ----------------------------------------
HUB_EN = 'https://www.bnm.md/en/content/supervised-entities-insurance-and-non-bank-lending'
HUB_RO = 'https://www.bnm.md/ro/content/entitati-supravegheate-asigurari-si-sectorul-nebancar'
BANKS  = 'https://www.bnm.md/en/content/authorized-banks-republic-moldova'

regdict = {
    'MD NBMO 1': BANKS,
    'MD NBMO 2': HUB_EN + '#art1',
    'MD NBMO 3': HUB_EN + '#art1',
    'MD NBMO 4': HUB_EN + '#art1',
    'MD NBMO 5': HUB_EN + '#art1',
}

topology = {
    regulatorName + ' 1': 'List of authorized banks of the Republic of Moldova',
    regulatorName + ' 2': 'List of authorized insurances of the Republic of Moldova',
    regulatorName + ' 3': 'List of Non-bank credit organizations of the Republic of Moldova',
    regulatorName + ' 4': 'List of Savings and Lending Associations of the Republic of Moldova',
    regulatorName + ' 5': 'List of Credit history bureaus of the Republic of Moldova',
}

# ListLabel: 1 = bank, 2 = insurance, 3 = bank & insurance, 4 = everything else
listlabel = {'1': 1, '2': 2, '3': 4, '4': 4, '5': 4}

# Registers are resolved by their LABEL on the hub page, never by file name:
# BNM re-uploads these files under new suffixes (_28, _14, (2)_2 ...) every few weeks.
LABELS = {
    '2a': 'Register of professional participants in the insurance market',
    '2b': 'List of insurance and/or reinsurance brokers',
    '2c': 'Register of insurance and bancassurance agents',
    '3':  'Register of authorized non-bank credit organizations',
    '4a': 'List of SLAs holding category A',
    '4b': 'List of SLAs holding category B',
    '5a': 'List of Credit history bureaus',
    '5b': 'List of entities (information sources)',
}
# Romanian fallback labels, used when the English page serves the wrong file.
LABELS_RO = {
    '2c': 'Registrul agenţilor de asigurare şi agenţilor bancassurance',
}

sqldict = {'bvdid': [], 'priority': [], 'ListLabel': [], 'Typology': [], 'EntryType': [], 'Name': [], 'InternalID_1': [], 'InternalID_1_type': [], 'InternalID_2': [],
           'InternalID_2_type': [], 'InternalID_3': [], 'InternalID_3_type': [], 'CoType': [], 'License_Type': [], 'Address_1': [], 'Address_2': [], 'City': [],
           'Zip': [], 'Cntry': [], 'Phone': [], 'Fax': [], 'Website': [], 'Email': [], 'RegulationType': [], 'RegulationTypeCode': [], 'RegulationDate': [], 'CancellationDate': [],
           'RegCtry': [], 'RegCode': [], 'ListCode': [], 'ListLanguage': [], 'ListValidityDate': [], 'ListName': [], 'ListProcessDate': [], 'LEI Code': [], 'BIC SWIFT Code': [], 'Name - Mother Company': [],
           'Address_1 - Mother company': [], 'Address_2 -  Mother company': [], 'City - Mother company': [], 'Zip - Mother company': [], 'Cntry - Mother company': [],
           'Phone - Mother company': []}

HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}
session = requests.Session()
session.headers.update(HEADERS)
session.verify = False

#------------------------------------------------ Begin_Function ----------------------------------------
def add_row(list_code, reg, validity='', **kw):
    """Append one record. Every column is written on every row, so the
    columns can never drift out of alignment (the old pad() could)."""
    base = {
        'ListLabel': listlabel[list_code],
        'RegCtry': 'MD',
        'RegCode': 'NBMO',
        'ListCode': list_code,
        'ListLanguage': 'EN',
        'ListName': topology[reg],
        'ListValidityDate': validity,
        'ListProcessDate': processdate,
        'RegulationType': 'Regulated',
        'Cntry': 'Moldova',
    }
    base.update({k: v for k, v in kw.items() if v not in (None, 'nan')})
    for col in sqldict:
        sqldict[col].append(base.get(col, ''))


def clean(val):
    """Normalise a cell: drop NaN, collapse newlines / repeated spaces."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ''
    txt = str(val).strip()
    if txt.lower() == 'nan':
        return ''
    return re.sub(r'\s+', ' ', txt.replace('\n', ' ')).strip()


def split_address(raw):
    """Return (address, zip) from a BNM address cell such as
    'MD-2012, mun. Chișinău, str. Alexandru cel Bun, 49'."""
    addr = clean(raw).split(';')[0].strip()
    m = re.search(r'MD\s?-?\s?\d{4}', addr)
    return addr, (m.group(0).replace(' ', '') if m else '')


def first_contact(raw, kind):
    """Pull a phone or an e-mail out of a mixed contact-details cell."""
    txt = clean(raw)
    if kind == 'email':
        m = re.search(r'[\w.\-+]+@[\w.\-]+\.\w+', txt)
        return m.group(0) if m else ''
    txt = re.sub(r'[\w.\-+]+@[\w.\-]+\.\w+', ' ', txt)          # strip e-mails first
    m = re.search(r'\(?\d[\d\s\-\(\)]{6,}\d', txt)
    return m.group(0).strip(' -') if m else ''


def find_validity(df, max_rows=6):
    """BNM stamps an 'Actualizat la / Updated at' date in the top rows."""
    block = df.head(max_rows)
    for _, row in block.iterrows():
        for cell in row:
            if isinstance(cell, (pd.Timestamp, datetime.datetime)):
                return cell.strftime('%Y-%m-%d')
    return ''


# ---------- link resolution (by label, never by file name) ----------
def load_hub(url):
    r = session.get(url, timeout=90)
    r.raise_for_status()
    return BeautifulSoup(r.text, 'html.parser'), r.url


def resolve_link(soup, base, label):
    """Find the download link whose anchor text starts with `label`."""
    best = None
    for a in soup.find_all('a', href=True):
        if not re.search(r'\.(xlsx|xls|docx|doc)\b', a['href'], re.I):
            continue
        text = re.sub(r'\s+', ' ', a.get_text(' ', strip=True))
        if label.lower() in text.lower():
            best = urljoin(base, a['href'])
            break
    return best


def download(url, tag):
    fn = unquote(url.split('/')[-1].split('?')[0])
    fp = os.path.join(tempfolder, fn)
    r = session.get(url, timeout=180)
    r.raise_for_status()
    with open(fp, 'wb') as fh:
        fh.write(r.content)
    print('[DL] {}: {} ({:,} bytes)'.format(tag, fn, len(r.content)))
    return fp


# ---------- legacy .doc reader: pure Python, no MS Word / pywin32 ----------
def _cfb_streams(data):
    """Minimal Compound-File-Binary reader -> {stream name: bytes}."""
    if data[:8] != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        raise ValueError('not an OLE2 compound file')
    ssz = 1 << struct.unpack_from('<H', data, 30)[0]          # sector size
    msz = 1 << struct.unpack_from('<H', data, 32)[0]          # mini-sector size
    n_fat = struct.unpack_from('<I', data, 44)[0]
    dir_start = struct.unpack_from('<I', data, 48)[0]
    mini_start = struct.unpack_from('<I', data, 60)[0]
    difat_start, n_difat = struct.unpack_from('<II', data, 68)

    def sector(i):
        off = 512 + i * ssz
        return data[off:off + ssz]

    fat_sectors = list(struct.unpack_from('<109I', data, 76))[:n_fat]
    nxt = difat_start                                          # follow the DIFAT chain
    while len(fat_sectors) < n_fat and nxt not in (0xFFFFFFFE, 0xFFFFFFFF):
        blk = sector(nxt)
        fat_sectors += list(struct.unpack_from('<%dI' % (ssz // 4 - 1), blk, 0))
        nxt = struct.unpack_from('<I', blk, ssz - 4)[0]
    fat = []
    for s in fat_sectors[:n_fat]:
        fat += list(struct.unpack_from('<%dI' % (ssz // 4), sector(s), 0))

    def chain(start, table):
        out, cur, guard = [], start, 0
        while cur < 0xFFFFFFFC and guard < 1 << 20:
            out.append(cur)
            cur = table[cur] if cur < len(table) else 0xFFFFFFFE
            guard += 1
        return out

    dir_bytes = b''.join(sector(s) for s in chain(dir_start, fat))
    entries = []
    for off in range(0, len(dir_bytes), 128):
        e = dir_bytes[off:off + 128]
        if len(e) < 128 or e[66] == 0:
            continue
        nlen = struct.unpack_from('<H', e, 64)[0]
        name = e[:max(nlen - 2, 0)].decode('utf-16-le', 'replace')
        start, size = struct.unpack_from('<I', e, 116)[0], struct.unpack_from('<I', e, 120)[0]
        entries.append((name, e[66], start, size))

    root = next((x for x in entries if x[1] == 5), None)
    mini_fat = []
    for s in chain(mini_start, fat):
        mini_fat += list(struct.unpack_from('<%dI' % (ssz // 4), sector(s), 0))
    mini_data = b''.join(sector(s) for s in chain(root[2], fat)) if root else b''

    streams = {}
    for name, typ, start, size in entries:
        if typ != 2:
            continue
        if size < 4096 and name != 'Root Entry':
            raw = b''.join(mini_data[i * msz:(i + 1) * msz] for i in chain(start, mini_fat))
        else:
            raw = b''.join(sector(s) for s in chain(start, fat))
        streams[name] = raw[:size]
    return streams


def doc_text(path):
    """Text stream of a Word 97-2003 .doc, with \\x07 cell markers preserved."""
    with open(path, 'rb') as fh:
        data = fh.read()
    st = _cfb_streams(data)
    wd = st['WordDocument']
    tbl_name = '1Table' if (struct.unpack_from('<H', wd, 0x0A)[0] & 0x0200) else '0Table'
    if tbl_name not in st:
        tbl_name = '0Table' if tbl_name == '1Table' else '1Table'
    tbl = st[tbl_name]

    fcClx, lcbClx = struct.unpack_from('<II', wd, 0x01A2)
    clx = tbl[fcClx:fcClx + lcbClx]
    i = 0
    while i < len(clx) and clx[i] == 0x01:                      # skip Prc blocks
        i += 3 + struct.unpack_from('<h', clx, i + 1)[0]
    if i >= len(clx) or clx[i] != 0x02:
        raise ValueError('piece table (PlcPcd) not found')
    lcb = struct.unpack_from('<I', clx, i + 1)[0]
    plc = clx[i + 5:i + 5 + lcb]
    n = (len(plc) - 4) // 12
    cps = list(struct.unpack_from('<%dI' % (n + 1), plc, 0))

    out = []
    for k in range(n):
        fc = struct.unpack_from('<I', plc, 4 * (n + 1) + 8 * k + 2)[0]
        compressed = bool(fc & 0x40000000)
        fc = (fc & 0x3FFFFFFF) // 2 if compressed else fc
        ln = cps[k + 1] - cps[k]
        raw = wd[fc:fc + (ln if compressed else ln * 2)]
        out.append(raw.decode('cp1251', 'replace') if compressed else raw.decode('utf-16-le', 'replace'))
    return ''.join(out)


def doc_table(path):
    """First table of a .doc / .docx as a DataFrame (no MS Word required)."""
    if path.lower().endswith('.docx'):
        doc = Document(path)
        if not doc.tables:
            return None
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        return pd.DataFrame(rows)

    text = doc_text(path)
    rows, cur, cell = [], [], []
    for ch in text:
        if ch == '\x07':
            if cur and not cell:            # second \x07 in a row = row terminator
                rows.append(cur)
                cur = []
            else:
                cur.append(''.join(cell).replace('\r', ' ').strip())
                cell = []
        elif ch in '\x0b\x0c':
            continue
        else:
            cell.append(ch)
    if cur:
        rows.append(cur)
    if not rows:
        return None
    width = max(len(r) for r in rows)
    return pd.DataFrame([r + [''] * (width - len(r)) for r in rows])


def is_header_row(row):
    joined = ' '.join(clean(c) for c in row).lower()
    return bool(re.search(r'denumire|nr\.?\s*d/o|state registration|название', joined))


# ======================== MAIN LOOP ========================
hub_soup, hub_url = load_hub(HUB_EN)
ro_soup = ro_url = None

for reg in regdict:
    print('Working with {}.'.format(reg))
    list_code = reg.split(' ')[-1]

    # ---- List 1: authorized banks (static HTML) ----
    if list_code == '1':
        r = session.get(BANKS, timeout=90)
        soup = BeautifulSoup(r.text.replace('<br>', '***'), 'html.parser')
        names = [e.text.strip() for e in soup.find_all('h3', {'class': 'bank-name'})
                 if len(e.text.lower().replace('banks', '').strip()) > 1]
        blocks = soup.find('div', {'class': 'container-post'}).find_all('div', {'class': 'bank-content'})[1:]
        for name, div in zip(names, blocks):
            vals = {}
            for line in div.find_all('div', {'class': 'line'}):
                dv = line.find_all('div')
                if len(dv) < 2:
                    continue
                lbl, val = dv[0].text, dv[1].text.strip()
                if 'Address' in lbl: vals['Address_1'], vals['Zip'] = split_address(val)
                if 'Phone'   in lbl: vals['Phone'] = val
                if 'Fax'     in lbl: vals['Fax'] = val
                if 'SWIFT'   in lbl: vals['BIC SWIFT Code'] = val
                if 'E-mail'  in lbl: vals['Email'] = val
                if 'WWW'     in lbl: vals['Website'] = val
            add_row(list_code, reg, Name=clean(name), CoType='Bank', **vals)

    # ---- List 2: insurers + brokers + bancassurance agents ----
    elif list_code == '2':
        # (a) Register of professional participants - FIRST table only (active insurers)
        fp = download(resolve_link(hub_soup, hub_url, LABELS['2a']), reg + ' insurers')
        raw = pd.read_excel(fp, sheet_name=0, header=None)
        validity = find_validity(raw)
        start = next(i for i in range(len(raw)) if is_header_row(raw.iloc[i])) + 1
        for i in range(start, len(raw)):
            if pd.isna(raw.iat[i, 0]):                          # blank row ends the 1st table
                break
            name = clean(raw.iat[i, 1])
            if not name:
                continue
            addr, zipc = split_address(raw.iat[i, 2])
            add_row(list_code, reg, validity=validity, Name=name,
                    InternalID_1=clean(raw.iat[i, 0]), InternalID_1_type='Registration number',
                    Address_1=addr, Zip=zipc,
                    Phone=first_contact(raw.iat[i, 3], 'phone'),
                    Email=first_contact(raw.iat[i, 3], 'email'),
                    CoType='Insurance company', License_Type=clean(raw.iat[i, 5])[:200])

        # (b) insurance / reinsurance brokers
        fp = download(resolve_link(hub_soup, hub_url, LABELS['2b']), reg + ' brokers')
        raw = pd.read_excel(fp, sheet_name=0, header=None)
        validity = find_validity(raw)
        for i in range(len(raw)):
            if pd.isna(pd.to_numeric(raw.iat[i, 0], errors='coerce')) or pd.isna(raw.iat[i, 1]):
                continue
            addr, zipc = split_address(raw.iat[i, 3])
            add_row(list_code, reg, validity=validity, Name=clean(raw.iat[i, 1]),
                    InternalID_1=clean(raw.iat[i, 2]), InternalID_1_type='IDNO',
                    Address_1=addr, Zip=zipc,
                    Phone=first_contact(raw.iat[i, 6], 'phone'),
                    Email=first_contact(raw.iat[i, 6], 'email'),
                    CoType='Insurance broker', License_Type=clean(raw.iat[i, 7])[:200])

        # (c) insurance & bancassurance agents.
        #     The English page currently links a stale COPY of the brokers register,
        #     so the downloaded workbook is validated and the RO page is used as fallback.
        url = resolve_link(hub_soup, hub_url, LABELS['2c'])
        fp = download(url, reg + ' agents')
        if 'Registrul' not in pd.ExcelFile(fp).sheet_names:
            print('[WARN] {}: EN page served the wrong file ({}) - falling back to the RO page'
                  .format(reg, os.path.basename(fp)))
            ro_soup, ro_url = load_hub(HUB_RO)
            fp = download(resolve_link(ro_soup, ro_url, LABELS_RO['2c']), reg + ' agents (RO)')
        raw = pd.read_excel(fp, sheet_name=0, header=None)
        validity = find_validity(raw)
        for i in range(len(raw)):
            if pd.isna(pd.to_numeric(raw.iat[i, 0], errors='coerce')) or pd.isna(raw.iat[i, 2]):
                continue
            addr, zipc = split_address(raw.iat[i, 4])
            add_row(list_code, reg, validity=validity, Name=clean(raw.iat[i, 2]),
                    InternalID_1=clean(raw.iat[i, 3]), InternalID_1_type='IDNO',
                    InternalID_2=clean(raw.iat[i, 1]),
                    InternalID_2_type='The unique registration code in the Register',
                    Address_1=addr, Zip=zipc,
                    Phone=first_contact(raw.iat[i, 5], 'phone'),
                    Email=first_contact(raw.iat[i, 6], 'email'),
                    CoType='Insurance agent')

    # ---- List 3: authorized non-bank credit organizations ----
    elif list_code == '3':
        fp = download(resolve_link(hub_soup, hub_url, LABELS['3']), reg)
        raw = pd.read_excel(fp, sheet_name=0, header=None)
        validity = find_validity(raw)
        for i in range(len(raw)):
            if pd.isna(raw.iat[i, 1]):
                continue
            name = clean(raw.iat[i, 1])
            state = clean(raw.iat[i, 11])
            # ticket: drop the orange rows / '(Radiată/ Excluded/ Исключена)' entries
            if not re.search(r'activ', state, re.I):
                continue
            if re.search(r'radiat|исключен|excluded', name, re.I):
                continue
            addr, zipc = split_address(raw.iat[i, 5])
            add_row(list_code, reg, validity=validity, Name=name,
                    InternalID_1=clean(raw.iat[i, 3]), InternalID_1_type='IDNO',
                    InternalID_2=clean(raw.iat[i, 0]), InternalID_2_type='Register number',
                    Address_1=addr, Zip=zipc,
                    Phone=first_contact(raw.iat[i, 6], 'phone'),
                    Email=first_contact(raw.iat[i, 6], 'email'),
                    RegulationDate=clean(raw.iat[i, 4])[:10],
                    CoType='Non-bank credit organization')

    # ---- List 4: savings & loan associations, category A + category B/ANC ----
    elif list_code == '4':
        for key in ('4a', '4b'):
            fp = download(resolve_link(hub_soup, hub_url, LABELS[key]), reg + ' ' + key)
            tbl = doc_table(fp)
            if tbl is None:
                print('[WARN] {}: no table found in {}'.format(reg, os.path.basename(fp)))
                continue
            print('[INFO] {} table: {}'.format(key, tbl.shape))
            for i in range(len(tbl)):
                row = tbl.iloc[i]
                if is_header_row(row):
                    continue
                name = clean(row.iloc[1])
                if not name or re.match(r'^\d+\.?$', name):
                    continue
                # ticket: from the category B list only, drop branches
                if key == '4b' and re.search(r'filiala|sucursala', name, re.I):
                    continue
                addr, zipc = split_address(row.iloc[3])
                add_row(list_code, reg, Name=name,
                        InternalID_1=clean(row.iloc[2]), InternalID_1_type='IDNO',
                        Address_1=addr, Zip=zipc,
                        Phone=first_contact(row.iloc[6], 'phone'),
                        License_Type=clean(row.iloc[4])[:200],
                        CoType='Savings and loan association')

    # ---- List 5: credit history bureaus + their information sources ----
    elif list_code == '5':
        # (a) the bureaus themselves - legacy .doc
        fp = download(resolve_link(hub_soup, hub_url, LABELS['5a']), reg + ' bureaus')
        tbl = doc_table(fp)
        print('[INFO] 5a table: {}'.format(None if tbl is None else tbl.shape))
        if tbl is not None:
            for i in range(len(tbl)):
                row = tbl.iloc[i]
                if is_header_row(row):
                    continue
                name = clean(row.iloc[1])
                if not name or re.match(r'^\d+\.?$', name):
                    continue
                addr, zipc = split_address(row.iloc[4])
                add_row(list_code, reg, Name=name,
                        InternalID_1=clean(row.iloc[2]), InternalID_1_type='IDNO',
                        Address_1=addr, Zip=zipc,
                        Phone=first_contact(row.iloc[5], 'phone'),
                        CoType='Credit history bureau')

        # (b) information sources holding contracts with the bureaus
        fp = download(resolve_link(hub_soup, hub_url, LABELS['5b']), reg + ' sources')
        raw = pd.read_excel(fp, sheet_name=0, header=None)
        validity = find_validity(raw)
        for i in range(len(raw)):
            if pd.isna(pd.to_numeric(raw.iat[i, 0], errors='coerce')) or pd.isna(raw.iat[i, 1]):
                continue
            name = clean(raw.iat[i, 1])
            # ticket: drop entities whose name contains 'radiata'
            if not name or re.search(r'radiat', name, re.I):
                continue
            add_row(list_code, reg, validity=validity, Name=name,
                    CoType='Credit history information source')

#------------------------------------------------ Begin_writer and save df to excel ----------------------------------------
os.chdir(scriptfolder)
df = pd.DataFrame(sqldict)
df = df[df['Name'] != '']

print('\nRows per list:')
print(df.groupby(['ListCode', 'ListName']).size().to_string())

df.to_excel(os.path.join(scriptfolder, filename), sheet_name='SQL Ready', index=False)
print('\nSaved {} rows to {}'.format(len(df), os.path.join(scriptfolder, filename)))
